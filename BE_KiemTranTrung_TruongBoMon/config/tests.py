from types import SimpleNamespace
from io import BytesIO
from unittest.mock import Mock, patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.urls import reverse
from docx import Document as WordDocument
from pypdf import PdfWriter
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIRequestFactory
from rest_framework_simplejwt.serializers import TokenObtainPairSerializer

from accounts.models import User
from approvals.models import TopicApproval
from topics.models import Topic
from config.api_views import (
    AdminOrDepartmentHeadOnly,
    AdminOrReadOnly,
    AdminUserViewSet,
    DocumentAccessPermission,
    DocumentSectionViewSet,
    DocumentViewSet,
    LecturerViewSet,
    SimilarityCheckViewSet,
    StudentViewSet,
    TechnologyViewSet,
    TopicDocumentViewSet,
    TopicFunctionViewSet,
    TopicTechnologyViewSet,
    TopicAccessPermission,
    TopicApprovalViewSet,
    TopicViewSet,
    has_full_api_access,
)
from config.file_title_extraction import (
    FileExtractionError,
    extract_document_text,
    find_title_candidates,
)
from config.serializers import (
    LoginTokenSerializer,
    TopicApprovalSerializer,
    TopicSerializer,
    StudentSerializer,
    UserSerializer,
)


class UserProfileDisplayTests(SimpleTestCase):
    def test_user_serializer_reads_display_fields_from_lecturer_profile(self):
        user = User(username="lecturer1", role=User.Role.LECTURER)
        user._state.fields_cache["lecturer_profile"] = SimpleNamespace(
            id=34,
            lecturer_code="GV001",
            full_name="Nguyen Van A",
            email="a@example.com",
            phone="0123456789",
            department_id=5,
            department=SimpleNamespace(pk=5, name="Khoa Công nghệ"),
            academic_degree="Thạc sĩ",
            specialization="Trí tuệ nhân tạo",
            is_active=True,
        )

        data = UserSerializer(user).data

        self.assertEqual(data["full_name"], "Nguyen Van A")
        self.assertEqual(data["email"], "a@example.com")
        self.assertEqual(data["phone"], "0123456789")

    def test_login_response_includes_linked_lecturer_profile(self):
        user = User(id=12, username="lecturer1", role=User.Role.LECTURER)
        user._state.fields_cache["lecturer_profile"] = SimpleNamespace(
            id=34,
            lecturer_code="GV001",
            full_name="Nguyen Van A",
            email="a@example.com",
            phone="0123456789",
            department_id=5,
            department=SimpleNamespace(pk=5, name="Khoa Công nghệ"),
            academic_degree="Thạc sĩ",
            specialization="Trí tuệ nhân tạo",
            is_active=True,
        )
        serializer = LoginTokenSerializer()
        serializer.user = user

        with patch.object(
            TokenObtainPairSerializer,
            "validate",
            return_value={"refresh": "refresh-token", "access": "access-token"},
        ):
            data = serializer.validate({"username": "lecturer1", "password": "secret"})

        self.assertEqual(data["access"], "access-token")
        self.assertEqual(data["user"]["id"], 12)
        self.assertEqual(data["user"]["full_name"], "Nguyen Van A")
        self.assertEqual(
            data["user"]["lecturer_profile"],
            {
                "id": 34,
                "lecturer_code": "GV001",
                "full_name": "Nguyen Van A",
                "email": "a@example.com",
                "phone": "0123456789",
                "department": 5,
                "department_name": "Khoa Công nghệ",
                "academic_degree": "Thạc sĩ",
                "specialization": "Trí tuệ nhân tạo",
                "is_active": True,
            },
        )

    def test_user_without_profile_uses_username_as_display_name(self):
        user = User(username="head1", role=User.Role.HEAD_OF_DEPARTMENT)

        self.assertEqual(user.full_name, "head1")
        self.assertIsNone(user.email)
        self.assertIsNone(user.phone)

    def test_department_head_uses_lecturer_profile_linked_to_account(self):
        user = User(id=12, username="head1", role=User.Role.HEAD_OF_DEPARTMENT)
        user._state.fields_cache["lecturer_profile"] = SimpleNamespace(
            id=34,
            lecturer_code="GV001",
            full_name="Nguyen Van A",
            email="a@example.com",
            phone="0123456789",
            department_id=5,
            department=SimpleNamespace(pk=5, name="Khoa Công nghệ"),
            academic_degree="Thạc sĩ",
            specialization="Trí tuệ nhân tạo",
            is_active=True,
        )

        data = UserSerializer(user).data

        self.assertEqual(data["full_name"], "Nguyen Van A")
        self.assertEqual(data["email"], "a@example.com")
        self.assertEqual(data["phone"], "0123456789")
        self.assertEqual(data["lecturer_profile"]["lecturer_code"], "GV001")
        self.assertEqual(data["lecturer_profile"]["department_name"], "Khoa Công nghệ")

    def test_student_user_reads_display_fields_from_student_profile(self):
        user = User(username="student1", role=User.Role.STUDENT)
        user._state.fields_cache["student_profile"] = SimpleNamespace(
            full_name="Tran Thi B",
            email="b@example.com",
            phone="0987654321",
        )

        self.assertEqual(user.full_name, "Tran Thi B")
        self.assertEqual(user.email, "b@example.com")
        self.assertEqual(user.phone, "0987654321")


class DepartmentHeadApiAccessTests(SimpleTestCase):
    def setUp(self):
        self.factory = APIRequestFactory()
        self.head = SimpleNamespace(
            is_authenticated=True,
            is_staff=False,
            role=User.Role.HEAD_OF_DEPARTMENT,
            id=1,
        )

    def test_department_head_has_full_api_access(self):
        self.assertTrue(has_full_api_access(self.head))

    def test_department_head_can_write_admin_protected_resources(self):
        request = self.factory.post("/api/v1/majors/")
        request.user = self.head

        self.assertTrue(AdminOrReadOnly().has_permission(request, view=None))
        self.assertTrue(
            AdminOrDepartmentHeadOnly().has_permission(request, view=None)
        )

    def test_department_head_can_manage_user_accounts(self):
        request = self.factory.patch("/api/v1/users/2/")
        request.user = self.head
        view = AdminUserViewSet()
        view.request = request

        self.assertTrue(
            all(permission.has_permission(request, view) for permission in view.get_permissions())
        )

    def test_department_head_can_edit_any_topic(self):
        request = self.factory.patch("/api/v1/topics/2/")
        request.user = self.head
        view = SimpleNamespace(action="update")
        topic = SimpleNamespace(created_by_id=99)
        permission = TopicAccessPermission()

        self.assertTrue(permission.has_permission(request, view))
        self.assertTrue(permission.has_object_permission(request, view, topic))

    def test_department_head_can_review_topics(self):
        request = self.factory.post("/api/v1/topics/2/review/")
        request.user = self.head
        view = SimpleNamespace(action="review")

        self.assertTrue(TopicAccessPermission().has_permission(request, view))

    def test_department_head_receives_unfiltered_department_data(self):
        request = self.factory.get("/api/v1/students/")
        request.user = self.head

        for viewset in (
            StudentViewSet,
            LecturerViewSet,
            TopicViewSet,
            TopicApprovalViewSet,
            SimilarityCheckViewSet,
            DocumentViewSet,
            DocumentSectionViewSet,
            TechnologyViewSet,
            TopicTechnologyViewSet,
            TopicFunctionViewSet,
            TopicDocumentViewSet,
        ):
            with self.subTest(viewset=viewset.__name__):
                view = viewset()
                view.request = request
                queryset = view.get_queryset()
                self.assertFalse(queryset.query.where.children)

    def test_department_head_has_document_object_access(self):
        request = self.factory.get("/api/v1/documents/2/")
        request.user = self.head
        document = SimpleNamespace(
            uploaded_by_id=99,
            topic=SimpleNamespace(
                created_by_id=98,
                lecturer_id=97,
                lecturer=SimpleNamespace(user_id=96),
            ),
        )

        self.assertTrue(
            DocumentAccessPermission().has_object_permission(
                request, view=None, obj=document
            )
        )


class TopicReviewPersistenceTests(SimpleTestCase):
    @patch("config.api_views.transaction.atomic")
    @patch("config.api_views.TopicApproval.objects.create")
    def test_review_persists_approval_timestamp(self, create_approval, _atomic):
        factory = APIRequestFactory()
        user = SimpleNamespace(id=8)
        topic = SimpleNamespace(
            id=3,
            status="SUBMITTED",
            updated_at=None,
            save=Mock(),
        )
        view = TopicViewSet()
        view.get_object = Mock(return_value=topic)
        view.get_serializer = Mock(
            return_value=SimpleNamespace(data={"id": 3, "status": "APPROVED"})
        )
        request = factory.post(
            "/api/v1/topics/3/review/",
            {"decision": TopicApproval.Action.APPROVE},
            format="json",
        )
        request.user = user
        view.action_map = {"post": "review"}
        request = view.initialize_request(request)
        request.user = user

        response = view.review(request, pk=3)

        self.assertEqual(response.data["status"], "APPROVED")
        approval = create_approval.call_args.kwargs
        self.assertEqual(approval["topic"], topic)
        self.assertEqual(approval["approver"], user)
        self.assertEqual(approval["action"], TopicApproval.Action.APPROVE)
        self.assertIsNotNone(approval["created_at"])
        topic.save.assert_called_once_with(update_fields=["status", "updated_at"])

    def test_approval_serializer_includes_topic_details(self):
        topic = Topic(id=3, topic_code="TC003", title="Đề tài thử nghiệm")
        approver = User(id=8, username="head1", role=User.Role.HEAD_OF_DEPARTMENT)
        approver._state.fields_cache["lecturer_profile"] = None
        approval = TopicApproval(
            id=5,
            topic=topic,
            approver=approver,
            action=TopicApproval.Action.REJECT,
            comment="Cần bổ sung phạm vi nghiên cứu.",
        )

        data = TopicApprovalSerializer(approval).data

        self.assertEqual(data["topic_title"], "Đề tài thử nghiệm")
        self.assertEqual(data["topic_code"], "TC003")
        self.assertEqual(data["approver_name"], "head1")
        self.assertEqual(data["comment"], "Cần bổ sung phạm vi nghiên cứu.")


class StudentProfileSerializerTests(SimpleTestCase):
    def test_student_profile_does_not_require_a_login_account(self):
        user_id = StudentSerializer().fields["user_id"]

        self.assertFalse(user_id.required)
        self.assertTrue(user_id.allow_null)


class TopicResourceApiTests(SimpleTestCase):
    def test_topic_resource_routes_are_registered(self):
        self.assertEqual(
            reverse("topic-resources", kwargs={"pk": 4}),
            "/api/v1/topics/4/resources/",
        )
        self.assertEqual(reverse("technology-list"), "/api/v1/technologies/")
        self.assertEqual(
            reverse("topic-technology-list"), "/api/v1/topic-technologies/"
        )
        self.assertEqual(reverse("topic-function-list"), "/api/v1/topic-functions/")
        self.assertEqual(reverse("topic-document-list"), "/api/v1/topic-documents/")
        self.assertEqual(
            reverse("similarity-check-extract-file"),
            "/api/v1/similarity/checks/extract-file/",
        )

    @patch("config.api_views.requests.post")
    @patch("config.api_views.TopicEmbedding.objects")
    def test_missing_approved_topic_embeddings_are_cached(
        self, embedding_manager, post
    ):
        embedding_manager.filter.return_value = []
        ai_response = Mock()
        ai_response.json.return_value = {
            "model_name": "tier1_best.pt",
            "embeddings": [[0.25, 0.75]],
        }
        post.return_value = ai_response

        result = SimilarityCheckViewSet()._load_candidate_embeddings(
            [{"id": 12, "title": "Đề tài kiểm thử"}]
        )

        self.assertEqual(result, {12: [0.25, 0.75]})
        post.assert_called_once()
        embedding_manager.update_or_create.assert_called_once()

    @patch("config.api_views.requests.post")
    @patch("config.api_views.TopicEmbedding.objects")
    def test_old_ai_service_without_embedding_route_uses_title_check_fallback(
        self, embedding_manager, post
    ):
        embedding_manager.filter.return_value = []
        missing_route = Mock(status_code=404)
        post.return_value = missing_route

        result = SimilarityCheckViewSet()._load_candidate_embeddings(
            [{"id": 12, "title": "Đề tài kiểm thử"}]
        )

        self.assertIsNone(result)
        post.assert_called_once()
        embedding_manager.update_or_create.assert_not_called()


class TopicAcademicYearValidationTests(SimpleTestCase):
    def test_semester_must_belong_to_selected_academic_year(self):
        serializer = TopicSerializer()
        academic_year = SimpleNamespace(id=1)
        semester = SimpleNamespace(academic_year_id=2)

        with self.assertRaises(ValidationError):
            serializer.validate(
                {"academic_year": academic_year, "semester": semester}
            )

    def test_matching_semester_and_academic_year_are_valid(self):
        serializer = TopicSerializer()
        academic_year = SimpleNamespace(id=1)
        semester = SimpleNamespace(academic_year_id=1)
        attributes = {"academic_year": academic_year, "semester": semester}

        self.assertEqual(serializer.validate(attributes), attributes)

    def test_topic_serializer_includes_only_active_student_assignments(self):
        serializer = TopicSerializer()
        assignments = [
            SimpleNamespace(
                status="ACTIVE",
                student_id=3,
                student=SimpleNamespace(full_name="Nguyen Van A", student_code="SV003"),
                lecturer_id=5,
                lecturer=SimpleNamespace(full_name="Tran Thi B"),
                assigned_at=None,
                academic_year_id=7,
                academic_year=SimpleNamespace(name="2025-2026"),
            ),
            SimpleNamespace(
                status="REMOVED",
                student_id=4,
                student=SimpleNamespace(full_name="Nguyen Van C", student_code="SV004"),
                lecturer_id=None,
                lecturer=None,
                assigned_at=None,
                academic_year_id=None,
                academic_year=None,
            ),
        ]
        topic = SimpleNamespace(
            _prefetched_objects_cache={"topic_students": assignments}
        )

        assignments = serializer.get_student_assignments(topic)
        self.assertEqual(len(assignments), 2)
        self.assertEqual(assignments[0]["academic_year_name"], "2025-2026")
        self.assertEqual(assignments[1]["status"], "REMOVED")
        self.assertEqual(
            serializer.get_active_students(topic),
            [
                {
                    "student_id": 3,
                    "student_name": "Nguyen Van A",
                    "student_code": "SV003",
                    "lecturer_id": 5,
                    "lecturer_name": "Tran Thi B",
                    "assigned_at": None,
                    "status": "ACTIVE",
                    "academic_year_id": 7,
                    "academic_year_name": "2025-2026",
                }
            ],
        )


class DocumentTitleExtractionTests(SimpleTestCase):
    def test_finds_labelled_title_and_ranks_it_first(self):
        text = (
            "TRƯỜNG ĐẠI HỌC KỸ THUẬT\n"
            "KHOA CÔNG NGHỆ\n"
            "TÊN ĐỀ TÀI:\n"
            "XÂY DỰNG HỆ THỐNG QUẢN LÝ THƯ VIỆN TRỰC TUYẾN\n"
            "Sinh viên thực hiện: Nguyễn Văn A\n"
        )

        candidates = find_title_candidates(text)

        self.assertTrue(candidates)
        self.assertEqual(
            candidates[0]["title"],
            "XÂY DỰNG HỆ THỐNG QUẢN LÝ THƯ VIỆN TRỰC TUYẾN",
        )
        self.assertGreaterEqual(candidates[0]["confidence"], 0.8)
        self.assertFalse(
            any(candidate["title"].startswith("TRƯỜNG") for candidate in candidates)
        )

    def test_extracts_text_from_docx_upload(self):
        document = WordDocument()
        document.add_paragraph("Tên đề tài:")
        document.add_paragraph("Phát triển ứng dụng quản lý học tập")
        content = BytesIO()
        document.save(content)
        upload = SimpleUploadedFile(
            "de-cuong.docx",
            content.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        extracted_text = extract_document_text(upload)

        self.assertIn("Tên đề tài:", extracted_text)
        self.assertIn("Phát triển ứng dụng quản lý học tập", extracted_text)

    def test_rejects_pdf_without_extractable_text(self):
        pdf_content = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf_content)
        upload = SimpleUploadedFile(
            "scan.pdf",
            pdf_content.getvalue(),
            content_type="application/pdf",
        )

        with self.assertRaisesMessage(
            FileExtractionError, "PDF scan/ảnh, cần OCR"
        ):
            extract_document_text(upload)
