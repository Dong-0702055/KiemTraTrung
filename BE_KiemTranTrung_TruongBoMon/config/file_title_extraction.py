import io
import re
import unicodedata
import zipfile

from docx import Document as WordDocument
from pypdf import PdfReader


MAX_UPLOAD_SIZE = 10 * 1024 * 1024
MAX_PDF_PAGES = 50
PREVIEW_CHARACTERS = 12_000
TITLE_LABEL = re.compile(
    r"^\s*(?:ten\s+de\s+tai|de\s+tai|"
    r"research\s+title|project\s+title|title|topic)\s*[:：\-]?\s*(.*)$",
    re.IGNORECASE,
)
IGNORED_LINE = re.compile(
    r"^(?:cong\s+hoa|doc\s+lap|truong\s+dai\s+hoc|"
    r"khoa\s+|bo\s+mon\s+|giang\s+vien\s+huong\s+dan|"
    r"sinh\s+vien\s+thuc\s+hien|ho\s+va\s+ten|"
    r"ma\s+so\s+sinh\s+vien|nam\s+hoc|khoa\s+luan|"
    r"bao\s+cao|luan\s+van|do\s+an\s+tot\s+nghiep)",
    re.IGNORECASE,
)
TOPIC_TERMS = re.compile(
    r"\b(xay\s*dung|phat\s*trien|nghien\s*cuu|ung\s*dung|thiet\s*ke|"
    r"he\s*thong|mo\s*hinh|giai\s*phap|phan\s*tich|quan\s*ly|"
    r"application|system|model|development|research)\b",
    re.IGNORECASE,
)


class FileExtractionError(ValueError):
    pass


def fold_diacritics(value):
    folded = []
    for character in value:
        if character in {"đ", "Đ"}:
            folded.append("d" if character == "đ" else "D")
            continue
        folded.append(
            "".join(
                part
                for part in unicodedata.normalize("NFD", character)
                if not unicodedata.combining(part)
            )
        )
    return "".join(folded)


def extract_document_text(uploaded_file):
    suffix = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""
    if suffix not in {"pdf", "docx"}:
        raise FileExtractionError("Chỉ hỗ trợ tệp PDF hoặc DOCX.")
    if uploaded_file.size > MAX_UPLOAD_SIZE:
        raise FileExtractionError("Tệp vượt quá giới hạn 10 MB.")

    content = uploaded_file.read()
    if suffix == "pdf":
        if not content.startswith(b"%PDF-"):
            raise FileExtractionError("Tệp không đúng định dạng PDF.")
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise FileExtractionError(
                    "PDF đang được bảo vệ bằng mật khẩu; hãy tải bản không khóa."
                )
            if len(reader.pages) > MAX_PDF_PAGES:
                raise FileExtractionError(
                    f"PDF vượt quá {MAX_PDF_PAGES} trang được phép xử lý."
                )
            text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )
        except FileExtractionError:
            raise
        except Exception as error:
            raise FileExtractionError(
                "Không đọc được PDF. Hãy kiểm tra tệp có bị lỗi hoặc mã hóa không."
            ) from error
    else:
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise FileExtractionError("Tệp không đúng định dạng DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded_size = sum(member.file_size for member in archive.infolist())
            if expanded_size > 30 * 1024 * 1024:
                raise FileExtractionError(
                    "DOCX giải nén vượt quá giới hạn xử lý 30 MB."
                )
            document = WordDocument(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend(
                    cell.text
                    for row in table.rows
                    for cell in row.cells
                )
            text = "\n".join(parts)
        except FileExtractionError:
            raise
        except Exception as error:
            raise FileExtractionError(
                "Không đọc được nội dung DOCX; hãy kiểm tra tệp có bị lỗi không."
            ) from error

    text = text.replace("\x00", "")
    if not text.strip():
        raise FileExtractionError(
            "Không trích xuất được chữ trong tệp. Nếu đây là PDF scan/ảnh, cần OCR "
            "hoặc tải bản PDF có thể chọn và sao chép văn bản."
        )
    return text[:PREVIEW_CHARACTERS]


def find_title_candidates(text, limit=5):
    lines = [
        re.sub(r"\s+", " ", line).strip(" \t•·-–—:：")
        for line in text.splitlines()
    ]
    lines = [line for line in lines if line]
    candidates = {}

    def add_candidate(value, position, labeled):
        value = re.sub(r"\s+", " ", value).strip(" \t•·-–—:：\"“”")
        folded_value = fold_diacritics(value)
        if not 12 <= len(value) <= 300 or IGNORED_LINE.search(folded_value):
            return
        alpha = [character for character in value if character.isalpha()]
        uppercase_ratio = (
            sum(character.isupper() for character in alpha) / len(alpha)
            if alpha
            else 0
        )
        score = 0.48 + (0.28 if labeled else 0)
        score += 0.12 if TOPIC_TERMS.search(folded_value) else 0
        score += 0.08 if uppercase_ratio >= 0.55 else 0
        score += max(0, 0.04 * (1 - position / max(len(lines), 1)))
        score = min(score, 0.99)
        current = candidates.get(value)
        if current is None or score > current["confidence"]:
            candidates[value] = {
                "title": value,
                "confidence": round(score, 2),
            }

    for index, line in enumerate(lines[:80]):
        label_match = TITLE_LABEL.match(fold_diacritics(line))
        if label_match:
            inline_title = line[label_match.start(1) :].strip()
            if inline_title:
                add_candidate(inline_title, index, True)
            else:
                for following in lines[index + 1 : index + 4]:
                    if following:
                        add_candidate(following, index, True)
                        break

        add_candidate(line, index, False)

    return sorted(
        candidates.values(),
        key=lambda candidate: (-candidate["confidence"], len(candidate["title"])),
    )[:limit]
